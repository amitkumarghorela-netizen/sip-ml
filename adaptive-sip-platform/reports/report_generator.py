"""
reports/report_generator.py
Report Engine (Phase 9): generates a professional multi-sheet Excel report
from backtest results (transactions + summary + comparison).
"""

from __future__ import annotations
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1):
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=start_row):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)


def generate_excel_report(compare_result: dict, output_path: str, asset_name: str = "ASSET"):
    """
    compare_result: output of core.backtest.compare_strategies()
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    # --- Comparison sheet ---
    ws_cmp = wb.create_sheet("Comparison")
    _write_df(ws_cmp, compare_result["comparison"])

    # --- One sheet per strategy: transactions + summary ---
    for key, res in compare_result["results"].items():
        sheet_name = res["strategy_name"][:28]  # excel sheet name limit
        ws = wb.create_sheet(sheet_name)

        txn_df = pd.DataFrame(res["transactions"])
        _write_df(ws, txn_df, start_row=1)

        # Summary block below transactions
        summary_start = len(txn_df) + 4
        ws.cell(row=summary_start, column=1, value="SUMMARY").font = Font(bold=True, size=12)
        for i, (k, v) in enumerate(res["summary"].items(), start=summary_start + 1):
            ws.cell(row=i, column=1, value=k).font = Font(bold=True)
            ws.cell(row=i, column=2, value=v)

        # Portfolio value chart
        if "Portfolio Value" in txn_df.columns and len(txn_df) > 1:
            chart = LineChart()
            chart.title = f"{sheet_name}: Portfolio Value vs Investment"
            chart.y_axis.title = "Value"
            chart.x_axis.title = "Period"
            pv_col_idx = list(txn_df.columns).index("Portfolio Value") + 1
            inv_col_idx = list(txn_df.columns).index("Total Investment") + 1
            data = Reference(ws, min_col=min(pv_col_idx, inv_col_idx),
                              max_col=max(pv_col_idx, inv_col_idx),
                              min_row=1, max_row=len(txn_df) + 1)
            chart.add_data(data, titles_from_data=True)
            ws.add_chart(chart, f"A{len(txn_df) + summary_start + 12}")

    wb.save(output_path)
    return output_path


def generate_csv_reports(compare_result: dict, output_dir: str):
    os.makedirs(output_dir, exist_ok=True)
    paths = []
    for key, res in compare_result["results"].items():
        df = pd.DataFrame(res["transactions"])
        p = os.path.join(output_dir, f"transactions_{key}.csv")
        df.to_csv(p, index=False)
        paths.append(p)
    cmp_path = os.path.join(output_dir, "comparison.csv")
    compare_result["comparison"].to_csv(cmp_path, index=False)
    paths.append(cmp_path)
    return paths
